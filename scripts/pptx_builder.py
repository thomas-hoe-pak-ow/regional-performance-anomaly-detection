import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

os.chdir(os.path.dirname(os.path.abspath(__file__)))

# ── 1. LOAD TEMPLATE ─────────────────────────────────────────────────────────
template_path = os.path.join('..', 'data', 'template.pptx')
prs = Presentation(template_path)

# ── 2. ACCESS SLIDE 1 ────────────────────────────────────────────────────────
slide = prs.slides[0]

# ── 3. INSERT CHART IMAGE ────────────────────────────────────────────────────
chart_path = os.path.join('..', 'data', 'loss_by_subcategory.png')

left   = Inches(0.5)   # distance from left edge
top    = Inches(1.5)   # distance from top edge
width  = Inches(9)     # image width
height = Inches(5)     # image height

slide.shapes.add_picture(chart_path, left, top, width, height)

# ── 4. ADD A SUBTITLE ────────────────────────────────────────────────────────
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(9), Inches(0.4))
tf = txBox.text_frame
tf.text = 'Source: Superstore Sales Dataset | Loss-making orders only'
tf.paragraphs[0].runs[0].font.size = Pt(9)
tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ── 5. SAVE OUTPUT ───────────────────────────────────────────────────────────
output_path = os.path.join('..', 'data', 'loss_report.pptx')
prs.save(output_path)
print(f"Presentation saved to {output_path}")