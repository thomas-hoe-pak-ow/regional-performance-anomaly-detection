import pandas as pd
import os
os.chdir(os.path.dirname(os.path.abspath(__file__)))

# Load the file
df = pd.read_csv("../data/Sample - Superstore.csv", encoding = "latin-1")
print("Number of rows: ", len(df))
print("Number of columns: ", len(df.columns))
print("Columns:", df.columns.to_list())
print("The columns are:", df.columns.to_list())
print("The shape of the dataset")
print(df.shape)

# Dataset: Structure and data tyes
# print("The structure and data types:")
# df.info()

# Dataset: Statistical summary
# print("The statistical summary:")
# print(df.describe(include = "all"))

# Dataset: Missing values - column-by-column missing count
# print("The column-by-column missing values:")
# print(df.isnull().sum())

# Dataset: A peek at few rows
# print("A look at 5 rows:")
# print(df.head()) # first 5 rows
# print(df.sample(5)) # 5 random rows

# From Superstore dataset, check for outlier by using interquartile range (IQR) method
def iqr_bounds(series, k=1.5):
    """
    Given a pandas Series of numbers, return (lower_bound, upper_bound)
    using the IQR method: anything outside these bounds is an outlier.
    """
    # Step 1: get Q1 (25th percentile) and Q3 (75th percentile) of `series`
    q1 = series.quantile(0.25)
    q3 = series.quantile(0.75)

    # Step 2: IQR = Q3 - Q1
    iqr = q3 -q1

    # Step 3: bounds = Q1 - k*IQR, Q3 + k*IQR
    lower = q1 - k*iqr
    upper = q3 + k*iqr

    return lower, upper

def flag_region_subcategory_outliers(df):
    """
    Aggregates Margin by Region + Sub-Category, then flags which combos
    are outliers using iqr_bounds().
    """
    # Step 1: groupby(['Region', 'Sub-Category'])[['Sales','Profit'], aggregate with .sum()
    # Use .reset_index() so Region/Sub-Category become columns again.
    grouped = df.groupby(['Region', 'Sub-Category'])[['Sales', 'Profit']].sum().reset_index()
    grouped['Margin'] = grouped['Profit'] / grouped['Sales']

    # Step 2: call iqr_bounds() on the aggregated Margin column
    lower, upper = iqr_bounds(grouped['Margin'])

    # Step 3: add an 'is_outlier' column: True where Margin is outside [lower, upper]
    grouped['is_outlier'] = (grouped['Margin'] < lower) | (grouped['Margin'] > upper)

    return grouped

# print(flag_region_subcategory_outliers(df))

# result = flag_region_subcategory_outliers(df)
# outliers = result[result['is_outlier']]

# print(outliers.sort_values('Margin'))
# print(len(outliers))

# lower, upper = iqr_bounds(result['Margin'])
# print(f"Bounds: {lower:.3f} to {upper:.3f}")
# print(f"Actual range: {result['Margin'].min():.3f} to {result['Margin'].max():.3f}")

# Results markdown
# Looking at margin by Region_Sub-Category, there is no outlier detected.
# This is likely because of the margin value being aggregated at Region_Sub-Category level, smoothing out the number in the process

# Next, we check individual rows and see if there are any with unusually high upper bound
def flag_group_outliers(group, col):
    """
    Given one Category's rows, flags which rows have an unusually high
    value in `col` (above that Category's own IQR upper bound).
    """
    # Step 1: get this group's bounds for `col`, using iqr_bounds()
    lower, upper = iqr_bounds(group[col])

    # Step 2: flag rows where col is ABOVE the upper bound
    # (we only care about "unusually HIGH", not low, for discount)
    group['is_outlier'] = group[col] > upper

    return group

def flag_discount_outliers(df):
    """
    Flags orders with unusually high discount relative to their Category's norm.
    """
    upper_by_category = df.groupby('Category')['Discount'].transform(
        lambda x: iqr_bounds(x)[1]
    )
    df['is_outlier'] = df['Discount'] > upper_by_category
    return df

result = flag_discount_outliers(df)
print(result.columns.tolist())      # confirm Category is present this time
print(result['is_outlier'].value_counts())
print(result[result['is_outlier']][['Category', 'Discount']].head(10))


import matplotlib.pyplot as plt

def plot_discount_by_category(df):
    """
    Boxplot of Discount distribution per Category, showing the IQR
    spread and outliers visually.
    """
    # Step 1: pandas DataFrames have a .boxplot() method that can group by a column.
    # Hint: df.boxplot(column=..., by=...) — which column do you want on the y-axis,
    # and which column defines the groups?
    df.boxplot(column='Discount', by='Category')

    # Step 2: give it a title and clean up the axis labels
    # Hint: plt.title(...), plt.suptitle('') to remove pandas' default subtitle,
    # plt.xlabel(...), plt.ylabel(...)
    plt.title('Boxplot of Discount by Category')
    plt.xlabel('Category')
    plt.ylabel('Discount')

    plt.savefig(os.path.join('..', 'data', 'discount_boxplot.png'), bbox_inches='tight')
    # Step 3: show it
    plt.show()

plot_discount_by_category(df)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

os.chdir(os.path.dirname(os.path.abspath(__file__)))

# ── 1. LOAD TEMPLATE ─────────────────────────────────────────────────────────
template_path = os.path.join('..', 'data', 'template.pptx')
prs = Presentation(template_path)

# ── 2. ACCESS SLIDE 1 ────────────────────────────────────────────────────────
slide = prs.slides[0]

# ── 3. ADD A TITLE ───────────────────────────────────────────────────────────
# (new — previous version relied on the template's own title, if any)
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf_title = title_box.text_frame
tf_title.text = 'Regional Performance Anomaly Detection — Key Findings'
tf_title.paragraphs[0].runs[0].font.size = Pt(24)
tf_title.paragraphs[0].runs[0].font.bold = True

# ── 4. INSERT CHART IMAGE ────────────────────────────────────────────────────
# (changed: new chart file, and shifted down/shrunk slightly to leave room
#  for the findings text below it)
chart_path = os.path.join('..', 'data', 'discount_boxplot.png')

left   = Inches(0.5)
top    = Inches(1.3)
width  = Inches(9)
height = Inches(4.2)

slide.shapes.add_picture(chart_path, left, top, width, height)

# ── 5. ADD FINDINGS SUMMARY ──────────────────────────────────────────────────
# (new — the actual "so what" for someone reading the slide)
findings_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.6), Inches(9), Inches(1.1))
tf_findings = findings_box.text_frame
tf_findings.word_wrap = True

findings_text = (
    "Q1 (Region \u00d7 Sub-Category margin): no statistical outliers detected \u2014 "
    "expected, since aggregation smooths extremes.\n"
    "Q2 (order-level discount): 703 of 9,994 orders (~7%) flagged as unusually high "
    "discount vs. their Category norm \u2014 concentrated in Office Supplies at 70\u201380% discount."
)
tf_findings.text = findings_text
for para in tf_findings.paragraphs:
    para.runs[0].font.size = Pt(12)

# ── 6. ADD A SUBTITLE / SOURCE LINE ─────────────────────────────────────────
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(9), Inches(0.4))
tf = txBox.text_frame
tf.text = 'Source: Superstore Sales Dataset | Method: IQR-based outlier detection'
tf.paragraphs[0].runs[0].font.size = Pt(9)
tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ── 7. SAVE OUTPUT ───────────────────────────────────────────────────────────
output_path = os.path.join('..', 'data', 'anomaly_detection_summary.pptx')
prs.save(output_path)
print(f"Presentation saved to {output_path}")